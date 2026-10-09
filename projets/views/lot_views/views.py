import json
from decimal import Decimal, InvalidOperation

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.http import JsonResponse
from django.shortcuts import get_object_or_404, redirect, render

from projets.decorators import chef_projet_required
from projets.exporters import ExcelExporter
from projets.manager import LigneHierarchique
from projets.models import LigneBordereau, LotProjet, Projet
from django.views.decorators.http import require_POST


@chef_projet_required
def saisie_bordereau(request, projet_id, lot_id):
    lot = get_object_or_404(LotProjet, id=lot_id, projet_id=projet_id)
    lot_root = lot.to_line_tree()

    data = [
        {
            'id': ligne.id,
            'numero': ligne.numero,
            'designation': ligne.designation,
            'unite': ligne.unite,
            'quantite': float(ligne.quantite),
            'prix_unitaire': float(ligne.pu),
            'montant': float(ligne.amount()),
            'niveau': ligne.level(),
            'est_titre': ligne.has_children(),
            'parent_id': ligne.parent.id if ligne.parent else None,
            '_expanded': False,
        }
        for ligne in lot_root.get_descendants()
    ]

    json_str = json.dumps(data, ensure_ascii=False)

    return render(request, 'projets/lots/saisie_bordereau.html', {
        'lot': lot,
        'root': lot_root,
        'lignes': json_str,
    })


def _iter_lignes_bordereau_hierarchiques(projet):
    """Retourne les lignes d'un projet dans le même ordre hiérarchique que la saisie du bordereau."""
    for lot in LotProjet.objects.filter(projet=projet).order_by('id'):
        lot_root = lot.to_line_tree()
        for ligne in lot_root.get_descendants():
            yield ligne


def _est_ligne_titre_bordereau(ligne):
    """Détermine si une ligne est une ligne de titre selon la logique du bordereau."""
    if hasattr(ligne, 'has_children') and ligne.has_children():
        return True
    return (
        not getattr(ligne, 'numero', None)
        or not getattr(ligne, 'unite', None)
        or getattr(ligne, 'quantite', None) is None
        or float(getattr(ligne, 'quantite', 0) or 0) == 0
    )


@chef_projet_required
def sauvegarder_lignes_bordereau(request, lot_id):
    if request.method == "POST":
        try:
            body = json.loads(request.body)
            lot = get_object_or_404(LotProjet, id=lot_id)

            lignes_existantes = {ligne.id: ligne for ligne in LigneBordereau.objects.filter(lot=lot)}
            id_mapping = {}
            lignes_existantes_utilisees = set()

            lignes = {}
            for index, row in enumerate(body):
                ligne_id = row.get('id')
                if ligne_id and ligne_id in lignes_existantes:
                    ligne = lignes_existantes[ligne_id]
                    ligne.numero = row.get('numero', '')
                    ligne.designation = row.get('designation', '')
                    ligne.unite = row.get('unite', '')
                    ligne.quantite = Decimal(str(row.get('quantite', 0)))
                    ligne.prix_unitaire = Decimal(str(row.get('prix_unitaire', 0)))
                    ligne.niveau = row.get('niveau', 0)
                    ligne.est_titre = row.get('est_titre', False)
                    ligne.ordre_affichage = index
                    ligne.montant_calcule = row.get('montant', 0)

                    lignes_existantes_utilisees.add(ligne_id)
                    lignes[ligne_id] = ligne
                    id_mapping[ligne_id] = ligne.id
                else:
                    ligne = LigneBordereau(
                        lot=lot,
                        numero=row.get('numero', ''),
                        designation=row.get('designation', ''),
                        unite=row.get('unite', ''),
                        quantite=Decimal(str(row.get('quantite', 0))),
                        prix_unitaire=Decimal(str(row.get('prix_unitaire', 0))),
                        niveau=row.get('niveau', 0),
                        est_titre=row.get('est_titre', False),
                        ordre_affichage=index,
                        montant_calcule=row.get('montant', 0)
                    )
                    lignes[ligne_id] = ligne

            try:
                for index, row in enumerate(body):
                    ligne_id = row.get('id')
                    parent_id = row.get('parent_id')
                    ligne: LigneBordereau = lignes[ligne_id]
                    if parent_id and parent_id in lignes:
                        ligne.parent = lignes[parent_id]
                    else:
                        ligne.parent = None
                    ligne.save()
                    id_mapping[ligne_id] = ligne.id

                lignes_a_supprimer = set(lignes_existantes.keys()) - lignes_existantes_utilisees
                if lignes_a_supprimer:
                    LigneBordereau.objects.filter(id__in=lignes_a_supprimer).delete()

                return JsonResponse({
                    'success': True,
                    'message': 'Lignes sauvegardées avec succès.',
                    'status': 'ok',
                    'lignes': id_mapping,
                }, status=200)
            except Exception as e:
                print(e)
                return JsonResponse({'error': str(e)}, status=400)

        except Exception as e:
            import traceback
            return JsonResponse({
                'status': 'error',
                'message': f"Erreur lors de la sauvegarde: {str(e)}",
                'traceback': traceback.format_exc()
            }, status=500)


def _parse_taux_tva(raw_value):
    """Valide le taux de TVA saisi ; vide => None (hérite du taux du projet)."""
    raw_value = (raw_value or '').strip()
    if raw_value == '':
        return None
    try:
        taux = Decimal(raw_value.replace(',', '.'))
    except (InvalidOperation, ValueError):
        raise ValueError("Le taux de TVA doit être un nombre.")
    if taux < 0 or taux > 100:
        raise ValueError("Le taux de TVA doit être compris entre 0 et 100.")
    return taux


def _preparer_contexte_bordereau(projet):
    lots = LotProjet.objects.filter(projet=projet).order_by('id')
    lots_data = []
    montant_total_ht = 0
    montant_total_tva = 0
    montant_total_ttc = 0
    total_lignes = 0

    for lot in lots:
        lignes = LigneBordereau.objects.filter(lot=lot).order_by('ordre_affichage')
        total_lot = sum((l.quantite or 0) * (l.prix_unitaire or 0) for l in lignes)
        if total_lot == 0:
            continue

        lines_root = LigneHierarchique({'id': 0, lot.nom: 'root'})
        lines_root.build_tree(lignes, lines_root)
        lignes_table = lines_root.export_to_table()

        lots_data.append({
            'lot': lot,
            'id': lot.id,
            'nom': lot.nom,
            'description': lot.description,
            'lignes_table': lignes_table,
            'total_lot': total_lot,
            'taux_tva': lot.taux_tva,
        })
        montant_total_ht += total_lot
        montant_total_tva += total_lot * (lot.taux_tva_applicable / 100)
        total_lignes += len(lignes_table)

    montant_total_ttc = montant_total_ht + montant_total_tva
    return {
        'projet': projet,
        'lots': lots_data,
        'montant_total_ht': montant_total_ht,
        'montant_total_tva': montant_total_tva,
        'montant_total_ttc': montant_total_ttc,
        'total_lots': len(lots_data),
        'total_lignes': total_lignes,
    }


@chef_projet_required
def lots_projet(request, projet_id):
    projet = get_object_or_404(Projet, id=projet_id)
    lots = LotProjet.objects.filter(projet=projet).order_by('id')

    ctx = {'projet': projet, 'lots': lots}

    if request.headers.get('HX-Request'):
        return render(request, 'projets/lots/_lots_content.html', ctx)

    return render(request, 'projets/lots/lots_projet.html', ctx)


@require_POST
@chef_projet_required
def ajouter_lot(request, projet_id):
    projet = get_object_or_404(Projet, id=projet_id)
    nom_lot = request.POST.get("nom", "").strip()
    is_xhr = request.headers.get('x-requested-with') == 'XMLHttpRequest'

    if not nom_lot:
        msg = "Le nom du lot ne peut pas être vide."
        if is_xhr:
            return JsonResponse({'success': False, 'message': msg}, status=400)
        messages.error(request, msg)
        return redirect('projets:lots_projet', projet_id=projet_id)

    try:
        taux_tva = _parse_taux_tva(request.POST.get('taux_tva', '20'))
    except ValueError as e:
        if is_xhr:
            return JsonResponse({'success': False, 'message': str(e)}, status=400)
        messages.error(request, str(e))
        return redirect('projets:lots_projet', projet_id=projet_id)

    lot = LotProjet.objects.create(
        projet=projet,
        nom=nom_lot,
        description=request.POST.get('description', '').strip(),
        taux_tva=taux_tva,
    )

    msg = f"Lot « {lot.nom} » ajouté avec succès."
    if is_xhr:
        return JsonResponse({'success': True, 'message': msg})

    messages.success(request, msg)
    return redirect('projets:lots_projet', projet_id=projet_id)


@require_POST
@chef_projet_required
def modifier_lot(request, projet_id, lot_id):
    lot = get_object_or_404(LotProjet, id=lot_id, projet_id=projet_id)
    is_xhr = request.headers.get('x-requested-with') == 'XMLHttpRequest'

    nouveau_nom = request.POST.get("nom", "").strip()

    if not nouveau_nom:
        msg = "Le nom du lot ne peut pas être vide."
        if is_xhr:
            return JsonResponse({'success': False, 'message': msg}, status=400)
        messages.error(request, msg)
        return redirect('projets:lots_projet', projet_id=projet_id)

    try:
        taux_tva = _parse_taux_tva(request.POST.get('taux_tva'))
    except ValueError as e:
        if is_xhr:
            return JsonResponse({'success': False, 'message': str(e)}, status=400)
        messages.error(request, str(e))
        return redirect('projets:lots_projet', projet_id=projet_id)

    lot.nom = nouveau_nom
    lot.description = request.POST.get("description", "").strip()
    lot.taux_tva = taux_tva
    lot.save()

    msg = f"Lot « {lot.nom} » modifié avec succès."
    if is_xhr:
        return JsonResponse({'success': True, 'message': msg})

    messages.success(request, msg)
    return redirect('projets:lots_projet', projet_id=projet_id)


@require_POST
@chef_projet_required
def supprimer_lot(request, projet_id, lot_id):
    lot = get_object_or_404(LotProjet, id=lot_id, projet_id=projet_id)
    nom = lot.nom
    is_xhr = request.headers.get('x-requested-with') == 'XMLHttpRequest'

    lot.delete()

    msg = f"Lot « {nom} » supprimé avec succès."
    if is_xhr:
        return JsonResponse({'success': True, 'message': msg})

    messages.success(request, msg)
    return redirect('projets:lots_projet', projet_id=projet_id)


@login_required
@chef_projet_required
def apercu_impression_bordereau(request, projet_id):
    """
    Aperçu avant impression du bordereau des prix.
    Réutilise la logique de lots_details mais avec un template dédié.
    """
    projet = get_object_or_404(Projet, id=projet_id)
    return render(request, 'projets/lots/apercu_impression_bordereau.html',
                  _preparer_contexte_bordereau(projet))

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from django.http import HttpResponse


@login_required
@chef_projet_required
def export_bordereau_excel(request, projet_id):
    projet = get_object_or_404(Projet, id=projet_id)
    contexte = _preparer_contexte_bordereau(projet)   # même contexte que l'aperçu
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Bordereau"

    # ---- Styles ----
    font_titre = Font(name='Calibri', size=14, bold=True, color='1F2023')
    font_entete = Font(name='Calibri', size=10, bold=True, color='FFFFFF')
    font_parent = Font(name='Calibri', size=10, bold=True, color='355B40')
    font_normal = Font(name='Calibri', size=10, color='1F2023')
    font_total = Font(name='Calibri', size=10, bold=True, color='FFFFFF')

    fill_entete = PatternFill('solid', fgColor='3A3D42')
    fill_lot = PatternFill('solid', fgColor='6D6E6F')
    fill_parent = PatternFill('solid', fgColor='EAEBEC')
    fill_total = PatternFill('solid', fgColor='355B40')

    align_left = Alignment(horizontal='left', vertical='center', wrap_text=True)
    align_right = Alignment(horizontal='right', vertical='center')
    align_center = Alignment(horizontal='center', vertical='center')
    align_indent_1 = Alignment(horizontal='left', vertical='center', indent=2, wrap_text=True)
    align_indent_2 = Alignment(horizontal='left', vertical='center', indent=4, wrap_text=True)
    align_indent_3 = Alignment(horizontal='left', vertical='center', indent=6, wrap_text=True)

    thin = Side(style='thin', color='D4D6D8')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    # ---- Titre ----
    ws.merge_cells('A1:F1')
    ws['A1'] = f"Bordereau des prix unitaires – {projet.nom}"
    ws['A1'].font = font_titre
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 24

    ws.merge_cells('A2:F2')
    ws['A2'] = f"Chantier : {projet.nom}"
    ws['A2'].alignment = Alignment(horizontal='center')
    ws['A2'].font = Font(size=10, italic=True, color='54585C')

    # ---- Colonnes ----
    headers = ['N°', 'Désignation', 'Unité', 'Quantité', 'PU (DH)', 'Montant (DH)']
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=col, value=header)
        cell.font = font_entete
        cell.fill = fill_entete
        cell.alignment = align_center if col in (1, 3) else (align_right if col >= 4 else align_left)
        cell.border = border

    row = 5
    for lot in contexte['lots']:
        # Bandeau du lot
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
        cell = ws.cell(row=row, column=1, value=f"LOT : {lot['nom']}")
        cell.font = font_total
        cell.fill = fill_lot
        cell.alignment = align_left
        for c in range(1, 7):
            ws.cell(row=row, column=c).border = border
        row += 1

        # Lignes du lot
        # print(f"Processing lot: {lot['nom']}", "Total lignes:", len(lot['lignes_table']))
        for ligne in lot['lignes_table']:
            # if not (hasattr(ligne, 'montant') and ligne.montant and ligne.montant > 0 and ligne.id != 0):
            #     continue
            ligne_id = ligne.get('id', 0)
            montant = ligne.get('montant') or 0

            # Ignorer la racine (id=0) et les lignes sans montant
            if ligne_id == 0:
                continue
            try:
                if float(montant) <= 0:
                    continue
            except (TypeError, ValueError):
                continue

            level = ligne.get('level', 0)
            indent = min(level, 3)
            align_des = [align_left, align_indent_1, align_indent_2, align_indent_3][indent]
            is_parent = ligne.get('is_parent', False)

            ws.cell(row=row, column=1, value=ligne.get('numero') or '').alignment = align_left
            ws.cell(row=row, column=2, value=ligne.get('designation') or '').alignment = align_des
            ws.cell(row=row, column=3, value=ligne.get('unite') or '').alignment = align_center
            ws.cell(row=row, column=4, value=float(ligne.get('quantite') or 0)).alignment = align_right
            ws.cell(row=row, column=5, value=float(ligne.get('prix_unitaire') or 0)).alignment = align_right
            ws.cell(row=row, column=6, value=float(montant)).alignment = align_right

            ws.cell(row=row, column=4).number_format = '#,##0.000'
            ws.cell(row=row, column=5).number_format = '#,##0.00'
            ws.cell(row=row, column=6).number_format = '#,##0.00'

            for c in range(1, 7):
                cell = ws.cell(row=row, column=c)
                cell.border = border
                cell.font = font_parent if is_parent else font_normal
                if is_parent:
                    cell.fill = fill_parent

            row += 1
        # Total du lot
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
        ws.cell(row=row, column=1, value=f"Total du lot {lot['nom']} HT :").font = font_total
        ws.cell(row=row, column=1).fill = fill_total
        ws.cell(row=row, column=1).alignment = align_right
        ws.cell(row=row, column=6, value=float(lot['total_lot'])).font = font_total
        ws.cell(row=row, column=6).fill = fill_total
        ws.cell(row=row, column=6).alignment = align_right
        ws.cell(row=row, column=6).number_format = '#,##0.00 "MAD"'
        for c in range(1, 7):
            ws.cell(row=row, column=c).border = border
        row += 2   # une ligne vide entre les lots

    # ---- Récapitulatif général ----
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
    ws.cell(row=row, column=1, value="RÉCAPITULATIF GÉNÉRAL").font = font_total
    ws.cell(row=row, column=1).fill = fill_total
    for c in range(1, 7):
        ws.cell(row=row, column=c).border = border
    row += 1

    recaps = [
        ("TOTAL GÉNÉRAL HT :", float(contexte['montant_total_ht'])),
        ("TOTAL GÉNÉRAL TVA :", float(contexte.get('montant_total_tva', 0))),
        ("TOTAL GÉNÉRAL TTC :", float(contexte['montant_total_ttc'])),
    ]
    for label, valeur in recaps:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
        ws.cell(row=row, column=1, value=label).font = font_total
        ws.cell(row=row, column=1).fill = fill_total
        ws.cell(row=row, column=1).alignment = align_right
        ws.cell(row=row, column=6, value=valeur).font = font_total
        ws.cell(row=row, column=6).fill = fill_total
        ws.cell(row=row, column=6).alignment = align_right
        ws.cell(row=row, column=6).number_format = '#,##0.00 "MAD"'
        for c in range(1, 7):
            ws.cell(row=row, column=c).border = border
        row += 1

    # ---- Largeurs de colonnes ----
    ws.column_dimensions['A'].width = 10
    ws.column_dimensions['B'].width = 55
    ws.column_dimensions['C'].width = 8
    ws.column_dimensions['D'].width = 12
    ws.column_dimensions['E'].width = 14
    ws.column_dimensions['F'].width = 18

    # ---- Freeze panes sur l'en-tête ----
    ws.freeze_panes = 'A5'

    # ---- Réponse ----
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    filename = f"Bordereau_{projet.nom}.xlsx".replace(' ', '_')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response

